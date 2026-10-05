from __future__ import annotations

import base64
import hashlib
import inspect
import io
import json
import logging
import os
import re
import time
import urllib.error
import urllib.request
import uuid
import xml.dom.minidom
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Literal, NotRequired, TypedDict, Union, cast

import numpy as np
import pandas as pd
from pypdf import PdfReader

from models.config import (
    DEFAULT_OLLAMA_URL,
    graph_extract_model,
    graph_extract_timeout_s,
    graph_force_rebuild,
    graph_max_source_chars_per_batch,
    graph_rag_enabled,
    rag_chunk_chars,
    rag_chunk_overlap,
    rag_embed_model,
    rag_max_files,
    success_benchmark_embed_cap,
    success_benchmark_top_k,
    vision_ingest_enabled,
    vision_ingest_max_image_bytes,
    vision_ingest_model,
    vision_ingest_timeout_s,
)

logger = logging.getLogger(__name__)

TEMPLATE_FILENAME = "uploaded_template.pptx"

# Path L: prompt for vision-LMM — output must be usable by a text-only strategist for PPTX.
_VISION_STRATEGIC_PROMPT = """You are assisting a strategic presentation workflow. The user cannot see this image.

Describe it in structured detail so another model can recreate the **ideas and layout** in PowerPoint slides.

Cover:
1) **Visual type** (e.g. whiteboard flowchart, hand-drawn diagram, chart, table photo, UI screenshot, slide photo).
2) **All readable text** — transcribe literally (labels, titles, axis text, sticky notes).
3) **Structure** — boxes/nodes, arrows, columns, hierarchy, numbering, grouping.
4) If it is a **chart or table**, describe axes, series, units, trends, and key numbers.
5) **Suggested slides** — bullet list: how many slides and what each should contain to mirror this content.

Use clear headings and bullet lists. Be factual; if something is illegible, say so."""


RASTER_VISION_EXTENSIONS = frozenset({"png", "jpg", "jpeg", "webp"})

SKIP_DIR_NAMES = frozenset({
    ".git",
    ".svn",
    "node_modules",
    "__pycache__",
    ".venv",
    "venv",
    ".mypy_cache",
    ".pytest_cache",
    ".ruff_cache",
    "dist",
    "build",
    "site-packages",
    ".idea",
    ".vscode",
    "mcps",
})

RAG_INGEST_EXTENSIONS = frozenset({
    "csv",
    "xlsx",
    "xlsm",
    "xls",
    "pdf",
    "json",
    "xml",
    "txt",
    "md",
    "pptx",
})


def _ollama_base_url() -> str:
    return os.environ.get("OLLAMA_HOST", DEFAULT_OLLAMA_URL).rstrip("/")


def _chunk_text(text: str, size: int, overlap: int) -> list[str]:
    text = (text or "").strip()
    if not text:
        return []
    if len(text) <= size:
        return [text]
    chunks: list[str] = []
    step = max(1, size - overlap)
    i = 0
    while i < len(text):
        chunks.append(text[i : i + size])
        i += step
    return chunks


def _ollama_embed_batches(base_url: str, model: str, texts: list[str]) -> list[list[float]] | None:
    if not texts:
        return []
    url = f"{base_url.rstrip('/')}/api/embed"
    batch_size = 32
    out: list[list[float]] = []
    for start in range(0, len(texts), batch_size):
        batch = texts[start : start + batch_size]
        payload: dict[str, Any] = {
            "model": model,
            "input": batch if len(batch) > 1 else batch[0],
        }
        try:
            req = urllib.request.Request(
                url,
                data=json.dumps(payload).encode("utf-8"),
                headers={"Content-Type": "application/json"},
                method="POST",
            )
            with urllib.request.urlopen(req, timeout=300) as resp:
                data = json.loads(resp.read().decode("utf-8"))
        except (urllib.error.URLError, urllib.error.HTTPError, json.JSONDecodeError, OSError) as e:
            logger.warning("ingestion.rag.embed_failed batch_start=%s err=%s", start, e)
            return None
        if "embeddings" in data:
            embs = data["embeddings"]
        elif "embedding" in data:
            embs = [data["embedding"]]
        else:
            logger.warning("ingestion.rag.embed_bad_shape keys=%s", list(data.keys()))
            return None
        if len(embs) != len(batch):
            logger.warning(
                "ingestion.rag.embed_count_mismatch expected=%s got=%s",
                len(batch),
                len(embs),
            )
            return None
        out.extend(embs)
    return out


def _cosine_rank(
    query_emb: list[float],
    doc_embs: list[list[float]],
) -> list[tuple[int, float]]:
    q = np.asarray(query_emb, dtype=np.float64)
    nq = np.linalg.norm(q) + 1e-12
    q = q / nq
    scored: list[tuple[int, float]] = []
    for i, d in enumerate(doc_embs):
        v = np.asarray(d, dtype=np.float64)
        nv = np.linalg.norm(v) + 1e-12
        v = v / nv
        scored.append((i, float(np.dot(q, v))))
    scored.sort(key=lambda x: x[1], reverse=True)
    return scored


def _lexical_rank(query: str, chunk_texts: list[str]) -> list[tuple[int, float]]:
    terms = [t for t in query.lower().split() if len(t) > 2]
    if not terms:
        return [(i, 1.0) for i in range(len(chunk_texts))]
    scores: list[tuple[int, float]] = []
    for i, chunk in enumerate(chunk_texts):
        tl = chunk.lower()
        s = float(sum(tl.count(w) for w in terms))
        scores.append((i, s))
    scores.sort(key=lambda x: x[1], reverse=True)
    return scores


def _cosine_similarity_vectors(a: list[float], b: list[float]) -> float:
    q = np.asarray(a, dtype=np.float64)
    v = np.asarray(b, dtype=np.float64)
    nq = np.linalg.norm(q) + 1e-12
    nv = np.linalg.norm(v) + 1e-12
    return float(np.dot(q / nq, v / nv))


def _cosine_to_confidence_percent(cos: float) -> int:
    """Map cosine in [-1, 1] to 0–100 for strategist-facing alignment."""
    x = max(-1.0, min(1.0, float(cos)))
    pct = (x + 1.0) / 2.0 * 100.0
    return int(round(max(0.0, min(100.0, pct))))


def _deck_digest_for_benchmark(
    user_prompt: str,
    slides: list[dict[str, Any]],
    *,
    max_chars: int = 14000,
) -> str:
    lines = [(user_prompt or "").strip(), "", "DECK_DIGEST:"]
    for i, s in enumerate(slides):
        if not isinstance(s, dict):
            continue
        t = str(s.get("title") or "").strip()
        b = str(s.get("body") or "").strip()
        chunk = b[:520] + ("…" if len(b) > 520 else "")
        lines.append(f"[{i + 1}] {t}\n{chunk}")
    return "\n".join(lines)[:max_chars]


def _entry_digest_for_benchmark(entry: dict[str, Any], *, max_chars: int = 14000) -> str:
    prompt = str(entry.get("prompt") or "").strip()
    deck = entry.get("deck") or {}
    slides = deck.get("slides") or []
    lines = [prompt, "", "DECK_DIGEST:"]
    for i, s in enumerate(slides):
        if not isinstance(s, dict):
            continue
        t = str(s.get("title") or "").strip()
        b = str(s.get("body") or "").strip()
        chunk = b[:520] + ("…" if len(b) > 520 else "")
        lines.append(f"[{i + 1}] {t}\n{chunk}")
    return "\n".join(lines)[:max_chars]


def _lexical_dice_similarity(a: str, b: str) -> float:
    wa = {t for t in re.split(r"\W+", (a or "").lower()) if len(t) > 2}
    wb = {t for t in re.split(r"\W+", (b or "").lower()) if len(t) > 2}
    if not wa or not wb:
        return 0.0
    inter = len(wa & wb)
    return float(2 * inter / (len(wa) + len(wb)))


def _vision_downscale_raster(data: bytes, max_side: int = 1568) -> bytes:
    """Reduce huge camera photos before base64; keeps dependency optional."""
    try:
        from PIL import Image
    except ImportError:
        return data
    try:
        im = Image.open(io.BytesIO(data))
        im = im.convert("RGB")
        w, h = im.size
        if max(w, h) <= max_side:
            return data
        scale = max_side / float(max(w, h))
        nw = max(1, int(w * scale))
        nh = max(1, int(h * scale))
        try:
            resample = Image.Resampling.LANCZOS
        except AttributeError:
            resample = Image.LANCZOS  # type: ignore[attr-defined]
        im = im.resize((nw, nh), resample)
        buf = io.BytesIO()
        im.save(buf, format="JPEG", quality=88, optimize=True)
        return buf.getvalue()
    except Exception as e:
        logger.warning("ingestion.vision_downscale_failed err=%s", e)
        return data


def _ollama_message_text(msg: dict[str, Any]) -> str:
    c = msg.get("content")
    if isinstance(c, str):
        return c.strip()
    if isinstance(c, list):
        parts: list[str] = []
        for p in c:
            if isinstance(p, dict):
                if p.get("type") == "text" and p.get("text"):
                    parts.append(str(p["text"]))
                elif p.get("text") is not None:
                    parts.append(str(p["text"]))
            elif isinstance(p, str):
                parts.append(p)
        return "\n".join(parts).strip()
    return ""


def _ollama_vision_describe(
    base_url: str,
    model: str,
    image_bytes: bytes,
    *,
    timeout: float,
) -> str | None:
    b64 = base64.standard_b64encode(image_bytes).decode("ascii")
    url = f"{base_url.rstrip('/')}/api/chat"
    payload: dict[str, Any] = {
        "model": model,
        "messages": [
            {
                "role": "user",
                "content": _VISION_STRATEGIC_PROMPT,
                "images": [b64],
            }
        ],
        "stream": False,
    }
    try:
        req = urllib.request.Request(
            url,
            data=json.dumps(payload).encode("utf-8"),
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            data = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        err_body = e.read().decode("utf-8", errors="replace")[:800]
        logger.warning(
            "ingestion.vision.http_error code=%s model=%r err=%s",
            e.code,
            model,
            err_body,
        )
        return None
    except (urllib.error.URLError, json.JSONDecodeError, OSError, TypeError) as e:
        logger.warning("ingestion.vision.request_failed model=%r err=%s", model, e)
        return None
    raw_msg = data.get("message")
    msg = raw_msg if isinstance(raw_msg, dict) else {}
    text = _ollama_message_text(msg)
    return text if text else None


def _pptx_plain_text(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    lines: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                lines.append(shape.text.strip())
    return "\n".join(lines).strip()


@dataclass(frozen=True)
class IngestionFile:
    """Explicit name + payload for ingestion (preferred from sync FastAPI handlers)."""

    name: str
    data: bytes


IngestionInput = Union[IngestionFile, tuple[str, bytes], Path, Any]


class IngestionMetadata(TypedDict, total=False):
    """Sidecar paths, template handle, and optional RAG trace."""

    template_metadata: str | None
    paths: list[str]
    rag_sources: NotRequired[list[str]]
    rag_mode: NotRequired[str]
    rag_files_scanned: NotRequired[int]
    rag_chunks_considered: NotRequired[int]
    graph_mode: NotRequired[str]
    graph_cache_hit: NotRequired[bool]
    graph_entities_total: NotRequired[int]
    graph_relations_total: NotRequired[int]
    graph_subgraph_entities: NotRequired[int]
    graph_subgraph_relations: NotRequired[int]


class IngestionResult(TypedDict):
    context_string: str
    metadata: IngestionMetadata


# --- Path X: Strategic knowledge graph (JSON-LD-style store) -----------------

KG_CACHE_FILENAME = "strategic_kg_cache.json"
_GRAPH_JSON_LD_CONTEXT = "https://relulayer.engine/ns/strategic-kg/v1"
_GRAPH_EXPORT_BASE_IRI = "https://relulayer.engine/kg/"


def _ttl_escape(s: str) -> str:
    return (
        (s or "")
        .replace("\\", "\\\\")
        .replace("\"", "\\\"")
        .replace("\n", "\\n")
        .replace("\r", "\\r")
        .replace("\t", "\\t")
    )


def _iri_fragment(s: str) -> str:
    """Best-effort safe fragment for Turtle local names."""
    raw = re.sub(r"[^A-Za-z0-9_\-]", "_", (s or "").strip())
    if not raw:
        return "x"
    if re.match(r"^[0-9]", raw):
        return "x_" + raw
    return raw[:96]


def _graph_document_to_turtle(document: dict[str, Any]) -> str:
    """
    Path GG: export Path X graph document to Turtle without rdflib.

    Nodes: entities as ``rl:Entity`` with labels/types/description/sources.
    Edges: direct predicate triple + a reified edge node with notes/sources.
    """
    ents = [e for e in (document.get("entities") or []) if isinstance(e, dict)]
    rels = [r for r in (document.get("relations") or []) if isinstance(r, dict)]
    lines: list[str] = [
        "@prefix rl: <https://relulayer.engine/ns/strategic-kg/v1#> .",
        "@prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .",
        "@prefix dct: <http://purl.org/dc/terms/> .",
        "",
    ]

    def ent_iri(e: dict[str, Any]) -> str:
        eid = str(e.get("id") or "").strip()
        frag = _iri_fragment(eid or f"{e.get('type','Other')}::{e.get('name','')}")
        return f"<{_GRAPH_EXPORT_BASE_IRI}{frag}>"

    for e in ents:
        iri = ent_iri(e)
        name = str(e.get("name") or "").strip()
        etype = str(e.get("type") or "Other").strip() or "Other"
        desc = str(e.get("description") or "").strip()
        srcs = e.get("sources") if isinstance(e.get("sources"), list) else []
        lines.append(f"{iri} a rl:Entity ;")
        if name:
            lines.append(f"  rdfs:label \"{_ttl_escape(name)}\" ;")
        lines.append(f"  rl:entityType \"{_ttl_escape(etype)}\" ;")
        if desc:
            lines.append(f"  dct:description \"{_ttl_escape(desc)}\" ;")
        for s in srcs:
            if s is None:
                continue
            lines.append(f"  rl:source \"{_ttl_escape(str(s))}\" ;")
        lines[-1] = lines[-1].rstrip(";") + " ."
        lines.append("")

    for i, r in enumerate(rels, start=1):
        subj = str(r.get("subject") or "").strip()
        obj = str(r.get("object") or "").strip()
        pred = str(r.get("predicate") or "Other").strip() or "Other"
        if not subj or not obj:
            continue
        subj_iri = f"<{_GRAPH_EXPORT_BASE_IRI}{_iri_fragment(subj)}>"
        obj_iri = f"<{_GRAPH_EXPORT_BASE_IRI}{_iri_fragment(obj)}>"
        pred_iri = f"rl:{_iri_fragment(pred.lower())}"
        lines.append(f"{subj_iri} {pred_iri} {obj_iri} .")

        bn = f"_:rel_{i}"
        lines.append(f"{bn} a rl:Relation ;")
        lines.append(f"  rl:subject {subj_iri} ;")
        lines.append(f"  rl:object {obj_iri} ;")
        lines.append(f"  rl:predicate \"{_ttl_escape(pred)}\" ;")
        st = str(r.get("subject_type") or "Other").strip() or "Other"
        ot = str(r.get("object_type") or "Other").strip() or "Other"
        lines.append(f"  rl:subjectType \"{_ttl_escape(st)}\" ;")
        lines.append(f"  rl:objectType \"{_ttl_escape(ot)}\" ;")
        note = str(r.get("notes") or "").strip()
        if note:
            lines.append(f"  rl:notes \"{_ttl_escape(note)}\" ;")
        srcs = r.get("sources") if isinstance(r.get("sources"), list) else []
        for s in srcs:
            if s is None:
                continue
            lines.append(f"  rl:source \"{_ttl_escape(str(s))}\" ;")
        lines[-1] = lines[-1].rstrip(";") + " ."
        lines.append("")

    return "\n".join(lines).strip() + "\n"

_GRAPH_EXTRACTION_SYSTEM = """You are a knowledge-graph extractor for enterprise strategy documents.
Output MUST be a single valid JSON object only — no markdown fences, no commentary.
Schema:
{"entities":[{"type":"Project|Metric|Stakeholder|Initiative|Risk|Constraint|Other","name":"string","description":"string"}],"relations":[{"subject":"string","subject_type":"same enum as entity types","predicate":"impacts|depends_on|owned_by|measures|delays|risk_to|mitigates|supports|blocks|targets|correlates_with|Other","object":"string","object_type":"same enum as entity types","notes":"string"}]}
Rules:
- Every entity name must be grounded in the SOURCE; do not invent proper nouns.
- Relations must be supported by the SOURCE; use notes to quote or paraphrase evidence.
- Prefer specific predicates; use Other only if no listed predicate fits.
- If nothing is extractable, return {"entities":[],"relations":[]}.
"""


def _strip_json_fence(s: str) -> str:
    s = (s or "").strip()
    m = re.match(r"^```(?:json)?\s*([\s\S]*?)```\s*$", s, re.IGNORECASE)
    if m:
        return m.group(1).strip()
    return s


def _ollama_chat_completion(
    base_url: str,
    model: str,
    system: str,
    user: str,
    *,
    timeout: float,
) -> str | None:
    url = f"{base_url.rstrip('/')}/api/chat"
    payload: dict[str, Any] = {
        "model": model,
        "messages": [
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
        "stream": False,
        "options": {"temperature": 0.15, "num_predict": 8192},
    }
    try:
        req = urllib.request.Request(
            url,
            data=json.dumps(payload).encode("utf-8"),
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            data = json.loads(resp.read().decode("utf-8"))
    except (urllib.error.HTTPError, urllib.error.URLError, json.JSONDecodeError, OSError, TypeError) as e:
        logger.warning("ingestion.graph.chat_failed model=%r err=%s", model, e)
        return None
    raw_msg = data.get("message")
    msg = raw_msg if isinstance(raw_msg, dict) else {}
    text = _ollama_message_text(msg)
    return text if text else None


def _parse_graph_json_object(raw: str) -> dict[str, Any] | None:
    cleaned = _strip_json_fence(raw)
    try:
        out = json.loads(cleaned)
    except json.JSONDecodeError:
        start = cleaned.find("{")
        end = cleaned.rfind("}")
        if start >= 0 and end > start:
            try:
                out = json.loads(cleaned[start : end + 1])
            except json.JSONDecodeError:
                logger.warning("ingestion.graph.json_parse_failed snippet=%r", cleaned[:240])
                return None
        else:
            return None
    return out if isinstance(out, dict) else None


def _stable_entity_id(entity_type: str, name: str) -> str:
    key = f"{entity_type.strip().lower()}::{name.strip().lower()}"
    return "e_" + hashlib.sha256(key.encode()).hexdigest()[:16]


def _entity_merge_key(e: dict[str, Any]) -> tuple[str, str] | None:
    t = str(e.get("type") or "Other").strip().lower()
    n = str(e.get("name") or "").strip()
    if not n:
        return None
    return (t, n.lower())


def _relation_merge_key(r: dict[str, Any]) -> tuple[str, str, str, str, str] | None:
    subj = str(r.get("subject") or "").strip().lower()
    obj = str(r.get("object") or "").strip().lower()
    pred = str(r.get("predicate") or "").strip().lower()
    st = str(r.get("subject_type") or "Other").strip().lower()
    ot = str(r.get("object_type") or "Other").strip().lower()
    if not subj or not obj:
        return None
    return (subj, st, pred, obj, ot)


def _merge_graph_batch(
    merged_entities: dict[tuple[str, str], dict[str, Any]],
    merged_relations: dict[tuple[str, str, str, str, str], dict[str, Any]],
    parsed: dict[str, Any],
    source_hint: str,
) -> None:
    for e in parsed.get("entities") or []:
        if not isinstance(e, dict):
            continue
        mk = _entity_merge_key(e)
        if not mk:
            continue
        name = str(e.get("name") or "").strip()
        etype = str(e.get("type") or "Other").strip() or "Other"
        desc = str(e.get("description") or "").strip()
        if mk in merged_entities:
            cur = merged_entities[mk]
            if len(desc) > len(str(cur.get("description") or "")):
                cur["description"] = desc
            srcs = cur.setdefault("sources", [])
            if source_hint and source_hint not in srcs:
                srcs.append(source_hint)
        else:
            merged_entities[mk] = {
                "id": _stable_entity_id(etype, name),
                "type": etype,
                "name": name,
                "description": desc,
                "sources": [source_hint] if source_hint else [],
            }

    for r in parsed.get("relations") or []:
        if not isinstance(r, dict):
            continue
        rk = _relation_merge_key(r)
        if not rk:
            continue
        if rk not in merged_relations:
            merged_relations[rk] = {
                "subject": str(r.get("subject") or "").strip(),
                "subject_type": str(r.get("subject_type") or "Other").strip() or "Other",
                "predicate": str(r.get("predicate") or "Other").strip() or "Other",
                "object": str(r.get("object") or "").strip(),
                "object_type": str(r.get("object_type") or "Other").strip() or "Other",
                "notes": str(r.get("notes") or "").strip(),
                "sources": [source_hint] if source_hint else [],
            }
        else:
            cur = merged_relations[rk]
            note = str(r.get("notes") or "").strip()
            if note and len(note) > len(str(cur.get("notes") or "")):
                cur["notes"] = note
            if source_hint:
                srcs = cur.setdefault("sources", [])
                if source_hint not in srcs:
                    srcs.append(source_hint)


def _ollama_extract_graph_for_text(
    base_url: str,
    model: str,
    source_text: str,
    *,
    timeout: float,
) -> dict[str, Any] | None:
    user = "SOURCE TEXT:\n\n" + source_text
    raw = _ollama_chat_completion(
        base_url,
        model,
        _GRAPH_EXTRACTION_SYSTEM,
        user,
        timeout=timeout,
    )
    if not raw:
        return None
    return _parse_graph_json_object(raw)


def _format_graph_subgraph_for_query(
    document: dict[str, Any],
    query: str,
    *,
    max_focus_entities: int = 12,
    max_extra_hop_entities: int = 10,
) -> tuple[str, int, int]:
    """Return markdown block and counts (entities shown, relations shown)."""
    entities_in = [e for e in (document.get("entities") or []) if isinstance(e, dict)]
    relations_in = [r for r in (document.get("relations") or []) if isinstance(r, dict)]
    q = (query or "").lower()
    terms = [t for t in re.split(r"\W+", q) if len(t) > 2]

    def score_entity(e: dict[str, Any]) -> float:
        name = str(e.get("name") or "")
        desc = str(e.get("description") or "")
        blob = f"{name} {desc}".lower()
        s = 0.0
        for t in terms:
            if t in name.lower():
                s += 3.0
            s += float(blob.count(t))
        return s

    scored = [(score_entity(e), e) for e in entities_in]
    scored.sort(key=lambda x: x[0], reverse=True)

    if any(s > 0 for s, _ in scored):
        focus = [e for s, e in scored if s > 0][:max_focus_entities]
    else:
        focus = [e for _, e in scored[: max(4, min(max_focus_entities, 8))]]

    name_by_lower: dict[str, str] = {}
    for e in entities_in:
        n = str(e.get("name") or "").strip()
        if n:
            name_by_lower.setdefault(n.lower(), n)

    focus_lower = {str(e.get("name") or "").strip().lower() for e in focus if e.get("name")}

    rel_pick: list[dict[str, Any]] = []
    extra_lower: set[str] = set()
    for r in relations_in:
        sn = str(r.get("subject") or "").strip()
        on = str(r.get("object") or "").strip()
        sl, ol = sn.lower(), on.lower()
        if sl in focus_lower or ol in focus_lower:
            rel_pick.append(r)
            if sl in focus_lower and ol:
                extra_lower.add(ol)
            if ol in focus_lower and sl:
                extra_lower.add(sl)

    extra_lower -= focus_lower
    hop_entities: list[dict[str, Any]] = []
    for el in list(extra_lower)[:max_extra_hop_entities]:
        canonical = name_by_lower.get(el)
        if not canonical:
            continue
        match = next(
            (x for x in entities_in if str(x.get("name") or "").strip().lower() == el),
            None,
        )
        if match:
            hop_entities.append(match)

    shown_entities = focus + hop_entities
    shown_lower = {str(e.get("name") or "").strip().lower() for e in shown_entities if e.get("name")}
    rel_filtered = [
        r
        for r in rel_pick
        if str(r.get("subject") or "").strip().lower() in shown_lower
        and str(r.get("object") or "").strip().lower() in shown_lower
    ]

    if not shown_entities and not rel_filtered:
        return "", 0, 0

    lines: list[str] = [
        "## Strategic knowledge graph (Path X — query-aligned subgraph)",
        "",
        "_Structured entities and relations extracted from the project folder; use for dependency "
        "chains (e.g. project delays impacting metrics or initiatives)._",
        "",
    ]
    lines.append("### Entities")
    lines.append("")
    for e in shown_entities:
        et = str(e.get("type") or "Other")
        nm = str(e.get("name") or "?")
        desc = str(e.get("description") or "").strip()
        srcs = e.get("sources") or []
        src_note = ""
        if isinstance(srcs, list) and srcs:
            src_note = f" _Sources: {', '.join(str(s) for s in srcs[:4])}"
            if len(srcs) > 4:
                src_note += "…"
        if desc:
            lines.append(f"- **{et}** — *{nm}*: {desc}{src_note}")
        else:
            lines.append(f"- **{et}** — *{nm}*{src_note}")
    lines.append("")
    lines.append("### Relations")
    lines.append("")
    if not rel_filtered:
        lines.append("_No relations matched the query-focused entity set; see RAG snippets for raw text._")
    else:
        for r in rel_filtered:
            subj = str(r.get("subject") or "?")
            pred = str(r.get("predicate") or "?")
            obj = str(r.get("object") or "?")
            note = str(r.get("notes") or "").strip()
            tail = f" — _{note}_" if note else ""
            lines.append(f"- **{subj}** — `{pred}` → **{obj}**{tail}")
    lines.append("")
    return "\n".join(lines), len(shown_entities), len(rel_filtered)


def _kg_cache_path(root: Path) -> Path:
    d = root / "assets" / "memory"
    d.mkdir(parents=True, exist_ok=True)
    return d / KG_CACHE_FILENAME


def _load_kg_cache(root: Path) -> dict[str, Any]:
    p = _kg_cache_path(root)
    if not p.is_file():
        return {"version": 1, "graphs": {}}
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {"version": 1, "graphs": {}}
    if not isinstance(data, dict):
        return {"version": 1, "graphs": {}}
    graphs = data.get("graphs")
    if not isinstance(graphs, dict):
        data["graphs"] = {}
    return data


def _save_kg_cache(root: Path, data: dict[str, Any]) -> None:
    p = _kg_cache_path(root)
    p.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def _extension(name: str) -> str:
    return Path(name).suffix.lower().lstrip(".")


def _normalize_upload(item: IngestionInput, index: int) -> tuple[str, bytes]:
    if isinstance(item, IngestionFile):
        return item.name, item.data
    if isinstance(item, tuple) and len(item) == 2:
        return str(item[0]), bytes(item[1])
    if isinstance(item, Path):
        path = item.resolve()
        return path.name, path.read_bytes()

    filename = getattr(item, "filename", None) or f"upload_{index}"
    read_fn = getattr(item, "read", None)
    if not callable(read_fn):
        raise TypeError(
            f"Unsupported file item at index {index}: expected IngestionFile, "
            "(name, bytes), Path, or object with read() -> bytes.",
        )
    raw = read_fn()
    if inspect.isawaitable(raw):
        raise TypeError(
            "Async file.read() is not supported inside IngestionEngine; await reads "
            "in the route and pass IngestionFile(name, data).",
        )
    return str(filename), bytes(cast(bytes, raw))


def _dataframe_markdown_preview(df: pd.DataFrame, max_rows: int = 10) -> str:
    sample = df.head(max_rows)
    columns = [str(c) for c in sample.columns]
    sep = "| " + " | ".join("---" for _ in columns) + " |"
    header = "| " + " | ".join(columns) + " |"
    rows = [
        "| " + " | ".join(str(sample.iloc[i, j]) for j in range(len(columns))) + " |"
        for i in range(len(sample))
    ]
    return "\n".join([header, sep, *rows])


def _tabular_body(df: pd.DataFrame) -> str:
    lines = [
        "**Columns & dtypes:**",
        "",
    ]
    for col in df.columns:
        lines.append(f"- `{col}` → `{df[col].dtype}`")
    lines.extend(["", f"**Row count:** {len(df)}", "", "**Markdown preview (first 10 rows):**", ""])
    lines.append(_dataframe_markdown_preview(df, 10))
    lines.append("")
    return "\n".join(lines)


class IngestionEngine:
    """
    Deterministic ingestion: tabular, documents, structured text, PPTX template capture,
    and raster images (PNG/JPG/WebP) via a local Ollama vision model (Path L).
    Produces a single context_string for Layer 3 DSPy modules plus metadata for tools.
    """

    def __init__(self, project_root: Path | None = None) -> None:
        root = (project_root or Path(__file__).resolve().parent.parent).resolve()
        self._project_root = root
        self._assets_dir = root / "assets"
        self._template_path = self._assets_dir / TEMPLATE_FILENAME

    def process_files(self, files: list[IngestionInput]) -> IngestionResult:
        logger.info("ingestion.start file_count=%s", len(files))
        parts: list[str] = []
        metadata: IngestionMetadata = {
            "template_metadata": None,
            "paths": [],
        }

        for idx, raw in enumerate(files):
            name, data = _normalize_upload(raw, idx)
            ext = _extension(name)
            size = len(data)
            logger.info(
                "ingestion.file index=%s name=%r ext=%r bytes=%s",
                idx,
                name,
                ext or "(none)",
                size,
            )

            try:
                if ext in ("csv",):
                    block = self._ingest_csv(name, data)
                elif ext in ("xlsx", "xlsm"):
                    block = self._ingest_excel(name, data)
                elif ext in ("xls",):
                    block = self._ingest_excel_legacy(name, data)
                elif ext in ("pdf",):
                    block = self._ingest_pdf(name, data)
                elif ext in ("json",):
                    block = self._ingest_json(name, data)
                elif ext in ("xml",):
                    block = self._ingest_xml(name, data)
                elif ext in ("pptx",):
                    block = self._ingest_pptx_template(name, data, metadata)
                elif ext in RASTER_VISION_EXTENSIONS:
                    block = self._ingest_image_raster(name, data)
                else:
                    logger.warning(
                        "ingestion.unknown_extension index=%s name=%r ext=%r treating_as_utf8_text",
                        idx,
                        name,
                        ext,
                    )
                    block = self._ingest_plain_text(name, data)

            except Exception:
                logger.exception(
                    "ingestion.failed index=%s name=%r ext=%r bytes=%s",
                    idx,
                    name,
                    ext,
                    size,
                )
                raise

            parts.append(block)
            logger.debug(
                "ingestion.block_done index=%s name=%r context_chars=%s",
                idx,
                name,
                len(block),
            )

        context_string = "\n\n---\n\n".join(parts) if parts else ""
        logger.info(
            "ingestion.complete context_chars=%s template=%r paths=%s",
            len(context_string),
            metadata.get("template_metadata"),
            metadata.get("paths", []),
        )
        return IngestionResult(context_string=context_string, metadata=metadata)

    def process_folder_rag(
        self,
        folder: str,
        query: str,
        *,
        top_k: int = 5,
        ollama_base_url: str | None = None,
        embed_model: str | None = None,
    ) -> IngestionResult:
        """
        Recursively scan a project folder, chunk extractable text, rank snippets against
        ``query`` (Strategic Directive) with Ollama embeddings; fall back to lexical overlap.
        """
        q = (query or "").strip()
        if not q:
            return IngestionResult(
                context_string="",
                metadata={
                    "template_metadata": None,
                    "paths": [],
                    "rag_mode": "skipped_empty_query",
                },
            )

        root = self._project_root.resolve()
        folder_path = self._resolve_safe_project_folder(folder, root)

        limit = rag_max_files()
        files = self._collect_rag_files(folder_path, limit)
        logger.info(
            "ingestion.rag.scan folder=%s max_files=%s matched=%s",
            folder_path,
            limit,
            len(files),
        )

        chunk_sz = rag_chunk_chars()
        overlap = rag_chunk_overlap()
        flat_chunks: list[str] = []
        chunk_refs: list[str] = []

        for fp in files:
            text = self._file_path_to_searchable_text(fp, root)
            if not text.strip():
                continue
            rel = fp.relative_to(root).as_posix()
            for i, ch in enumerate(_chunk_text(text, chunk_sz, overlap)):
                flat_chunks.append(ch)
                chunk_refs.append(f"{rel}#{i}")

        if not flat_chunks:
            return IngestionResult(
                context_string="",
                metadata={
                    "template_metadata": None,
                    "paths": [],
                    "rag_mode": "no_text_extracted",
                    "rag_files_scanned": len(files),
                    "rag_chunks_considered": 0,
                },
            )

        model = (embed_model or rag_embed_model()).strip()
        base = (ollama_base_url or _ollama_base_url()).rstrip("/")

        q_emb = _ollama_embed_batches(base, model, [q])
        doc_emb: list[list[float]] | None = None
        if q_emb is not None and len(q_emb) == 1:
            doc_emb = _ollama_embed_batches(base, model, flat_chunks)

        if doc_emb is not None and len(doc_emb) == len(flat_chunks):
            ranked = _cosine_rank(q_emb[0], doc_emb)
            mode = "ollama"
        else:
            ranked = _lexical_rank(q, flat_chunks)
            mode = "lexical"
            logger.info(
                "ingestion.rag.using_lexical chunks=%s (embeddings unavailable or mismatched)",
                len(flat_chunks),
            )

        top_idx = [idx for idx, _score in ranked[: max(1, top_k)]]
        parts: list[str] = []
        seen_paths: set[str] = set()
        for rank_order, idx in enumerate(top_idx):
            ref = chunk_refs[idx]
            src_path, _, _chunk_i = ref.partition("#")
            seen_paths.add(src_path)
            parts.append(
                f"### RAG rank {rank_order + 1} — `{src_path}`\n\n{flat_chunks[idx]}",
            )

        header = (
            "## Retrieved context (multi-document RAG, top "
            f"{len(parts)} of {len(flat_chunks)} snippets, mode={mode})\n\n---\n\n"
        )
        context_string = header + "\n\n---\n\n".join(parts)

        metadata: IngestionMetadata = {
            "template_metadata": None,
            "paths": sorted(seen_paths),
            "rag_sources": sorted(seen_paths),
            "rag_mode": mode,
            "rag_files_scanned": len(files),
            "rag_chunks_considered": len(flat_chunks),
        }
        logger.info(
            "ingestion.rag.complete mode=%s top_k=%s chunks=%s",
            mode,
            len(parts),
            len(flat_chunks),
        )
        return IngestionResult(context_string=context_string, metadata=metadata)

    def _folder_content_fingerprint(self, folder_path: Path, root: Path) -> str:
        files = self._collect_rag_files(folder_path, rag_max_files())
        parts: list[str] = []
        for fp in sorted(files, key=lambda x: str(x)):
            try:
                st = fp.stat()
                rel = fp.relative_to(root).as_posix()
                parts.append(f"{rel}:{st.st_mtime_ns}:{st.st_size}")
            except (OSError, ValueError):
                continue
        blob = "\n".join(parts)
        return hashlib.sha256(blob.encode()).hexdigest()

    def process_folder_graph_context(
        self,
        folder: str,
        query: str,
        *,
        ollama_base_url: str | None = None,
    ) -> IngestionResult:
        """
        Path X: build or load a strategic knowledge graph from the same files as folder RAG,
        then inject a query-aligned subgraph into GTM context (entities + relations).
        """
        if not graph_rag_enabled():
            return IngestionResult(
                context_string="",
                metadata={
                    "template_metadata": None,
                    "paths": [],
                    "graph_mode": "disabled",
                },
            )

        q = (query or "").strip()
        root = self._project_root.resolve()
        folder_path = self._resolve_safe_project_folder(folder, root)
        folder_key = folder_path.relative_to(root).as_posix()

        fingerprint = self._folder_content_fingerprint(folder_path, root)
        cache = _load_kg_cache(root)
        graphs = cache.setdefault("graphs", {})
        assert isinstance(graphs, dict)

        cache_hit = False
        document: dict[str, Any] | None = None
        entry = graphs.get(folder_key)
        if (
            not graph_force_rebuild()
            and isinstance(entry, dict)
            and entry.get("fingerprint") == fingerprint
            and isinstance(entry.get("document"), dict)
        ):
            document = cast(dict[str, Any], entry["document"])
            cache_hit = True
            logger.info(
                "ingestion.graph.cache_hit folder=%s entities=%s",
                folder_key,
                len(document.get("entities") or []),
            )

        if document is None:
            limit = rag_max_files()
            files = self._collect_rag_files(folder_path, limit)
            batch_chars = graph_max_source_chars_per_batch()
            batches: list[tuple[str, str]] = []
            buf_parts: list[str] = []
            buf_size = 0

            for fp in files:
                text = self._file_path_to_searchable_text(fp, root)
                if not text.strip():
                    continue
                rel = fp.relative_to(root).as_posix()
                block = f"### FILE: {rel}\n{text}\n"
                if buf_size + len(block) > batch_chars and buf_parts:
                    batches.append((f"batch_{len(batches)}", "\n".join(buf_parts)))
                    buf_parts = [block]
                    buf_size = len(block)
                else:
                    buf_parts.append(block)
                    buf_size += len(block)
            if buf_parts:
                batches.append((f"batch_{len(batches)}", "\n".join(buf_parts)))

            merged_ent: dict[tuple[str, str], dict[str, Any]] = {}
            merged_rel: dict[tuple[str, str, str, str, str], dict[str, Any]] = {}
            base = (ollama_base_url or _ollama_base_url()).rstrip("/")
            model = graph_extract_model()
            timeout = graph_extract_timeout_s()

            for hint, corpus in batches:
                parsed = _ollama_extract_graph_for_text(
                    base,
                    model,
                    corpus,
                    timeout=timeout,
                )
                if parsed:
                    _merge_graph_batch(merged_ent, merged_rel, parsed, hint)
                else:
                    logger.warning("ingestion.graph.batch_empty hint=%s", hint)

            document = {
                "@context": _GRAPH_JSON_LD_CONTEXT,
                "folder": folder_key,
                "entities": list(merged_ent.values()),
                "relations": list(merged_rel.values()),
            }
            graphs[folder_key] = {
                "fingerprint": fingerprint,
                "updated_at": time.time(),
                "document": document,
            }
            try:
                _save_kg_cache(root, cache)
            except OSError as e:
                logger.warning("ingestion.graph.cache_save_failed err=%s", e)

            logger.info(
                "ingestion.graph.built folder=%s batches=%s entities=%s relations=%s",
                folder_key,
                len(batches),
                len(merged_ent),
                len(merged_rel),
            )

        assert document is not None
        block, n_ent, n_rel = _format_graph_subgraph_for_query(document, q)
        meta: IngestionMetadata = {
            "template_metadata": None,
            "paths": [folder_key],
            "graph_mode": "cache_hit" if cache_hit else "rebuilt",
            "graph_cache_hit": cache_hit,
            "graph_entities_total": len(document.get("entities") or []),
            "graph_relations_total": len(document.get("relations") or []),
            "graph_subgraph_entities": n_ent,
            "graph_subgraph_relations": n_rel,
        }
        return IngestionResult(context_string=block, metadata=meta)

    def export_folder_knowledge_graph(
        self,
        folder: str,
        *,
        export_format: Literal["jsonld", "turtle"] = "jsonld",
        force_rebuild: bool = False,
        ollama_base_url: str | None = None,
    ) -> tuple[str, bytes, str]:
        """
        Path GG: export the Path X strategic knowledge graph for a project folder.

        Returns (filename, file_bytes, mime_type).

        - ``jsonld``: exports the cached graph document (JSON-LD-ish dict with @context/entities/relations).
        - ``turtle``: exports an RDF Turtle representation (lightweight; no rdflib dependency).

        If the graph isn't cached (or the folder changed), this will build it (same pipeline as
        :meth:`process_folder_graph_context`) before exporting.
        """
        root = self._project_root.resolve()
        folder_path = self._resolve_safe_project_folder(folder, root)
        folder_key = folder_path.relative_to(root).as_posix()

        fingerprint = self._folder_content_fingerprint(folder_path, root)
        cache = _load_kg_cache(root)
        graphs = cache.setdefault("graphs", {})
        assert isinstance(graphs, dict)
        entry = graphs.get(folder_key)

        needs_build = True
        if (
            not force_rebuild
            and isinstance(entry, dict)
            and entry.get("fingerprint") == fingerprint
            and isinstance(entry.get("document"), dict)
        ):
            needs_build = False

        if needs_build:
            # Build+cache via the existing Path X pipeline (query doesn't matter for full export).
            _ = self.process_folder_graph_context(
                folder_key,
                query="export",
                ollama_base_url=ollama_base_url,
            )
            cache = _load_kg_cache(root)
            graphs = cache.get("graphs") if isinstance(cache.get("graphs"), dict) else {}
            entry = graphs.get(folder_key) if isinstance(graphs, dict) else None

        doc = entry.get("document") if isinstance(entry, dict) else None
        if not isinstance(doc, dict):
            raise ValueError(
                f"No knowledge graph cached for folder {folder_key!r}. "
                "Enable graph RAG and/or run a draft with project_folder first.",
            )

        if export_format == "jsonld":
            name = f"strategic_kg_{_iri_fragment(folder_key)}.jsonld"
            blob = json.dumps(doc, ensure_ascii=False, indent=2).encode("utf-8")
            return name, blob, "application/ld+json"
        if export_format == "turtle":
            name = f"strategic_kg_{_iri_fragment(folder_key)}.ttl"
            blob = _graph_document_to_turtle(doc).encode("utf-8")
            return name, blob, "text/turtle"
        raise ValueError(f"Unsupported export_format: {export_format!r}")

    def _resolve_safe_project_folder(self, folder: str, root: Path) -> Path:
        p = Path(folder).expanduser()
        resolved = p.resolve() if p.is_absolute() else (root / p).resolve()
        try:
            resolved.relative_to(root)
        except ValueError as e:
            raise ValueError(
                f"Project folder must stay inside project root {root}; got {resolved}",
            ) from e
        if not resolved.is_dir():
            raise ValueError(f"Project folder is not a directory: {resolved}")
        return resolved

    def _collect_rag_files(self, folder: Path, limit: int) -> list[Path]:
        out: list[Path] = []
        for path in sorted(folder.rglob("*")):
            if len(out) >= limit:
                break
            if not path.is_file():
                continue
            if any(part in SKIP_DIR_NAMES for part in path.parts):
                continue
            ext = path.suffix.lower().lstrip(".")
            if ext not in RAG_INGEST_EXTENSIONS:
                continue
            try:
                rel = path.relative_to(self._project_root.resolve())
            except ValueError:
                continue
            parts = rel.parts
            if len(parts) >= 2 and parts[0] == "assets" and parts[1] == "outputs":
                continue
            out.append(path)
        return out

    def _file_path_to_searchable_text(self, path: Path, root: Path) -> str:
        rel = path.relative_to(root).as_posix()
        try:
            data = path.read_bytes()
        except OSError as e:
            logger.warning("ingestion.rag.read_failed path=%r err=%s", path, e)
            return ""
        ext = path.suffix.lower().lstrip(".")
        try:
            return self._raw_bytes_to_searchable_text(rel, data, ext)
        except Exception:
            logger.exception("ingestion.rag.extract_failed path=%r", path)
            return ""

    def _raw_bytes_to_searchable_text(self, rel_name: str, data: bytes, ext: str) -> str:
        header = f"[{rel_name}]\n"
        if ext == "csv":
            for encoding in ("utf-8", "utf-8-sig", "latin-1"):
                try:
                    df = pd.read_csv(io.BytesIO(data), encoding=encoding)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                df = pd.read_csv(io.BytesIO(data), encoding="utf-8", errors="replace")
            body = df.head(300).to_string(index=False)
            return header + body

        if ext in ("xlsx", "xlsm"):
            bio = io.BytesIO(data)
            xl = pd.ExcelFile(bio)
            df = pd.read_excel(xl, sheet_name=0)
            body = df.head(300).to_string(index=False)
            return header + f"sheets: {xl.sheet_names}\n" + body

        if ext == "xls":
            try:
                df = pd.read_excel(io.BytesIO(data), engine="xlrd")
            except ImportError:
                return header + "(xls requires xlrd)"
            return header + df.head(300).to_string(index=False)

        if ext == "pdf":
            reader = PdfReader(io.BytesIO(data))
            texts = [(page.extract_text() or "") for page in reader.pages]
            return header + "\n\n".join(texts).strip()

        if ext == "json":
            text = data.decode("utf-8", errors="replace")
            parsed = json.loads(text)
            compact = json.dumps(parsed, ensure_ascii=False)[:80000]
            return header + compact

        if ext == "xml":
            text = data.decode("utf-8", errors="replace")
            pretty = xml.dom.minidom.parseString(text).toprettyxml(indent="  ")
            return header + pretty.strip()[:80000]

        if ext == "pptx":
            return header + _pptx_plain_text(data)

        if ext in ("txt", "md"):
            return header + data.decode("utf-8", errors="replace")

        return header + data.decode("utf-8", errors="replace")

    def _ingest_csv(self, name: str, data: bytes) -> str:
        logger.info("ingestion.csv name=%r", name)
        for encoding in ("utf-8", "utf-8-sig", "latin-1"):
            try:
                df = pd.read_csv(io.BytesIO(data), encoding=encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            df = pd.read_csv(io.BytesIO(data), encoding="utf-8", errors="replace")
        return f"### Tabular (CSV): `{name}`\n\n" + _tabular_body(df)

    def _ingest_excel(self, name: str, data: bytes) -> str:
        logger.info("ingestion.xlsx name=%r", name)
        bio = io.BytesIO(data)
        xl = pd.ExcelFile(bio)
        sheet_names = xl.sheet_names
        logger.info(
            "ingestion.excel_sheets name=%r sheets=%s",
            name,
            sheet_names,
        )
        df = pd.read_excel(xl, sheet_name=0)
        header = (
            f"### Tabular (Excel): `{name}`\n\n"
            f"**Sheets:** {', '.join(sheet_names)}\n\n"
            f"**Active sheet:** `{sheet_names[0]}`\n\n"
        )
        return header + _tabular_body(df)

    def _ingest_excel_legacy(self, name: str, data: bytes) -> str:
        logger.info("ingestion.xls name=%r", name)
        try:
            df = pd.read_excel(io.BytesIO(data), engine="xlrd")
        except ImportError:
            logger.error("ingestion.xls_missing_xlrd name=%r", name)
            raise RuntimeError(
                "Reading .xls requires the optional `xlrd` package; use .xlsx or install xlrd.",
            ) from None
        return f"### Tabular (Excel .xls): `{name}`\n\n" + _tabular_body(df)

    def _ingest_pdf(self, name: str, data: bytes) -> str:
        logger.info("ingestion.pdf name=%r", name)
        reader = PdfReader(io.BytesIO(data))
        page_count = len(reader.pages)
        texts: list[str] = []
        for i, page in enumerate(reader.pages):
            t = page.extract_text() or ""
            texts.append(t)
            logger.debug("ingestion.pdf_page name=%r page=%s chars=%s", name, i + 1, len(t))
        body = "\n\n".join(texts).strip()
        logger.info(
            "ingestion.pdf_done name=%r pages=%s total_chars=%s",
            name,
            page_count,
            len(body),
        )
        return f"### PDF: `{name}`\n\n**Pages:** {page_count}\n\n{body}"

    def _ingest_json(self, name: str, data: bytes) -> str:
        logger.info("ingestion.json name=%r", name)
        text = data.decode("utf-8", errors="replace")
        parsed = json.loads(text)
        formatted = json.dumps(parsed, indent=2, ensure_ascii=False)
        logger.debug("ingestion.json_done name=%r chars=%s", name, len(formatted))
        return f"### JSON: `{name}`\n\n```json\n{formatted}\n```"

    def _ingest_xml(self, name: str, data: bytes) -> str:
        logger.info("ingestion.xml name=%r", name)
        text = data.decode("utf-8", errors="replace")
        pretty = xml.dom.minidom.parseString(text).toprettyxml(indent="  ")
        logger.debug("ingestion.xml_done name=%r chars=%s", name, len(pretty))
        return f"### XML: `{name}`\n\n```xml\n{pretty.strip()}\n```"

    def _ingest_pptx_template(
        self,
        name: str,
        data: bytes,
        metadata: IngestionMetadata,
    ) -> str:
        logger.info("ingestion.pptx_template name=%r bytes=%s", name, len(data))
        self._assets_dir.mkdir(parents=True, exist_ok=True)
        self._template_path.write_bytes(data)
        path_str = str(self._template_path.resolve())
        metadata["template_metadata"] = path_str
        metadata["paths"].append(path_str)
        logger.info(
            "ingestion.pptx_saved path=%r (from upload name=%r)",
            path_str,
            name,
        )
        return (
            f"### PPTX template\n\n"
            f"A presentation template was uploaded as `{name}` and stored for deterministic "
            f"compilation at:\n\n`{path_str}`\n\n"
            f"(Binary slide masters are not inlined into the text context.)"
        )

    def _ingest_image_raster(self, name: str, data: bytes) -> str:
        """Path L: describe image with Ollama vision LLM for GTM / PPTX structure replication."""
        logger.info("ingestion.image name=%r bytes=%s", name, len(data))
        header = f"### Vision-ingested image: `{name}`\n\n"
        note = (
            "_This block was generated by a local vision model so the strategist can mirror "
            "diagrams, whiteboards, or slide photos in the final deck._\n\n"
        )
        mx = vision_ingest_max_image_bytes()
        if len(data) > mx:
            return (
                header
                + note
                + f"**Skipped vision call:** file is {len(data)} bytes (limit {mx} bytes). "
                "Increase `RELULAYER_VISION_MAX_MB` or use a smaller image.\n"
            )
        if not vision_ingest_enabled():
            return (
                header
                + note
                + "**Vision ingestion disabled** (`RELULAYER_VISION_INGEST=0`). "
                "Set to `1` and configure `RELULAYER_VISION_MODEL` with an Ollama vision-capable model.\n"
            )
        base = _ollama_base_url()
        model = vision_ingest_model()
        to_send = _vision_downscale_raster(data)
        desc = _ollama_vision_describe(
            base,
            model,
            to_send,
            timeout=vision_ingest_timeout_s(),
        )
        if not desc:
            return (
                header
                + note
                + f"**Vision call failed or returned empty** (model `{model}` at `{base}`). "
                "Run `ollama pull` for that tag and check server logs.\n"
            )
        logger.info(
            "ingestion.image_vision_done name=%r model=%r desc_chars=%s",
            name,
            model,
            len(desc),
        )
        return (
            header
            + note
            + "## Model description (replicate structure in slides)\n\n"
            + desc.strip()
            + "\n"
        )

    def _ingest_plain_text(self, name: str, data: bytes) -> str:
        text = data.decode("utf-8", errors="replace")
        logger.info("ingestion.text name=%r chars=%s", name, len(text))
        return f"### Unstructured text: `{name}`\n\n{text}"


# --- Strategic memory: successful materializations (Layer 4) -------------------


def _success_store_enabled() -> bool:
    return os.environ.get("RELULAYER_SUCCESS_STORE", "1").strip().lower() not in (
        "0",
        "false",
        "no",
        "off",
    )


def _success_store_max_entries() -> int:
    return max(10, int(os.environ.get("RELULAYER_SUCCESS_STORE_MAX", "250")))


def _success_few_shot_k() -> int:
    return max(0, min(8, int(os.environ.get("RELULAYER_SUCCESS_FEW_SHOT_K", "3"))))


class SuccessStore:
    """
    Append-only JSONL under ``assets/memory/success_store.jsonl`` with optional
    Ollama embeddings over the strategic directive for similarity retrieval.
    """

    def __init__(self, project_root: Path | None = None) -> None:
        root = (project_root or Path(__file__).resolve().parent.parent).resolve()
        self._root = root
        self._dir = root / "assets" / "memory"
        self._path = self._dir / "success_store.jsonl"

    def _ensure_dir(self) -> None:
        self._dir.mkdir(parents=True, exist_ok=True)

    def load_entries(self) -> list[dict[str, Any]]:
        if not self._path.is_file():
            return []
        raw = self._path.read_text(encoding="utf-8", errors="replace")
        out: list[dict[str, Any]] = []
        for line in raw.splitlines():
            line = line.strip()
            if not line:
                continue
            try:
                out.append(json.loads(line))
            except json.JSONDecodeError:
                logger.warning("success_store.skip_bad_json_line")
        return out

    def _write_all(self, entries: list[dict[str, Any]]) -> None:
        self._ensure_dir()
        text = "\n".join(json.dumps(e, ensure_ascii=False) for e in entries)
        if text:
            text += "\n"
        self._path.write_text(text, encoding="utf-8")

    def _maybe_trim(self, entries: list[dict[str, Any]]) -> None:
        mx = _success_store_max_entries()
        if len(entries) <= mx:
            return
        entries.sort(key=lambda e: float(e.get("ts", 0.0)))
        self._write_all(entries[-mx:])

    def record_materialization(
        self,
        *,
        user_prompt: str,
        persona: str,
        slides: list[dict[str, Any]],
        design_specs: dict[str, Any] | None,
        ollama_base_url: str,
    ) -> bool:
        if not _success_store_enabled():
            logger.info("success_store.record_skipped disabled")
            return False
        prompt = (user_prompt or "").strip()
        if not prompt or not slides:
            return False

        base = ollama_base_url.rstrip("/")
        model = rag_embed_model()
        emb: list[float] | None = None
        qe = _ollama_embed_batches(base, model, [prompt])
        if qe and len(qe) == 1:
            emb = qe[0]

        entry: dict[str, Any] = {
            "id": str(uuid.uuid4()),
            "ts": time.time(),
            "persona": (persona or "").strip() or "govtech",
            "prompt": prompt,
            "embedding": emb,
            "deck": {
                "slides": slides,
                "design_specs": design_specs,
            },
        }
        self._ensure_dir()
        with self._path.open("a", encoding="utf-8") as f:
            f.write(json.dumps(entry, ensure_ascii=False) + "\n")
        logger.info(
            "success_store.recorded id=%s slides=%s embedded=%s",
            entry["id"][:8],
            len(slides),
            emb is not None,
        )
        self._maybe_trim(self.load_entries())
        return True

    def compute_historical_benchmark(
        self,
        *,
        user_prompt: str,
        draft_slides: list[dict[str, Any]],
        persona: str,
        ollama_base_url: str,
    ) -> dict[str, Any]:
        """
        Path AA: embed the current directive + deck digest and compare to historical successes
        in ``success_store.jsonl`` (top-k cosine mean → confidence percent + best-match label).
        """
        base_out: dict[str, Any] = {
            "confidence_percent": 0,
            "mode": "disabled",
            "reference_label": "",
            "mean_top_k_similarity": None,
            "top_matches": [],
            "display_hint": "",
        }
        if not _success_store_enabled():
            return base_out

        entries = self.load_entries()
        if not entries:
            base_out["mode"] = "insufficient_history"
            base_out["display_hint"] = (
                "No rows in success_store.jsonl yet — materialize a winning deck and record success "
                "to populate historical benchmarks."
            )
            return base_out

        k = success_benchmark_top_k()
        cap = success_benchmark_embed_cap()
        draft_digest = _deck_digest_for_benchmark(user_prompt, draft_slides)

        subset = list(entries)
        if len(subset) > cap:
            titles = " ".join(
                str(s.get("title") or "") for s in draft_slides if isinstance(s, dict)
            )
            q_lex = f"{(user_prompt or '').strip()}\n{titles}"
            prompts = [str(e.get("prompt") or "") for e in subset]
            pick = [idx for idx, _ in _lexical_rank(q_lex, prompts)[:cap]]
            subset = [subset[i] for i in pick]

        entry_digests = [_entry_digest_for_benchmark(e) for e in subset]
        base = ollama_base_url.rstrip("/")
        model = rag_embed_model()
        batch: list[str] = [draft_digest] + entry_digests
        embs = _ollama_embed_batches(base, model, batch)

        top_matches: list[dict[str, Any]] = []
        reference_label = ""
        mean_top: float | None = None
        mode = "unavailable"
        confidence_percent = 0

        if embs is not None and len(embs) == len(batch):
            d0 = embs[0]
            scored: list[tuple[float, dict[str, Any]]] = []
            for i, ent in enumerate(subset):
                sim = _cosine_similarity_vectors(d0, embs[1 + i])
                if (ent.get("persona") or "") == (persona or ""):
                    sim = min(1.0, sim + 0.04)
                scored.append((sim, ent))
            scored.sort(key=lambda x: -x[0])
            top = scored[:k]
            mean_top = sum(t[0] for t in top) / len(top) if top else None
            confidence_percent = (
                _cosine_to_confidence_percent(mean_top) if mean_top is not None else 0
            )
            mode = "embedding"
            if top:
                best_prompt = str(top[0][1].get("prompt") or "").replace("\n", " ").strip()
                reference_label = (
                    (best_prompt[:140] + ("…" if len(best_prompt) > 140 else ""))
                    if best_prompt
                    else "(untitled directive)"
                )
            for sim, ent in top:
                p = str(ent.get("prompt") or "").replace("\n", " ").strip()
                eid = str(ent.get("id") or "")
                top_matches.append(
                    {
                        "prompt_excerpt": p[:200] + ("…" if len(p) > 200 else ""),
                        "similarity": round(max(-1.0, min(1.0, sim)), 4),
                        "entry_id_prefix": eid[:8],
                        "persona": str(ent.get("persona") or ""),
                    },
                )
        else:
            mode = "lexical_fallback"
            scored_lex: list[tuple[float, dict[str, Any]]] = []
            for ent in subset:
                d = _entry_digest_for_benchmark(ent)
                sim = _lexical_dice_similarity(draft_digest, d)
                scored_lex.append((sim, ent))
            scored_lex.sort(key=lambda x: -x[0])
            top = scored_lex[:k]
            mean_top = sum(t[0] for t in top) / len(top) if top else None
            confidence_percent = int(
                round(max(0.0, min(100.0, (mean_top or 0.0) * 100.0))),
            )
            if top:
                best_prompt = str(top[0][1].get("prompt") or "").replace("\n", " ").strip()
                reference_label = (
                    (best_prompt[:140] + ("…" if len(best_prompt) > 140 else ""))
                    if best_prompt
                    else ""
                )
            for sim, ent in top:
                p = str(ent.get("prompt") or "").replace("\n", " ").strip()
                eid = str(ent.get("id") or "")
                top_matches.append(
                    {
                        "prompt_excerpt": p[:200] + ("…" if len(p) > 200 else ""),
                        "similarity": round(max(0.0, min(1.0, sim)), 4),
                        "entry_id_prefix": eid[:8],
                        "persona": str(ent.get("persona") or ""),
                    },
                )

        if reference_label:
            display_hint = (
                f"{confidence_percent}% alignment with a successful historical project: {reference_label}"
            )
        else:
            display_hint = f"{confidence_percent}% historical confidence vs top prior materializations (Path AA)"

        logger.info(
            "success_store.benchmark mode=%s confidence=%s k=%s entries=%s",
            mode,
            confidence_percent,
            k,
            len(subset),
        )
        return {
            "confidence_percent": confidence_percent,
            "mode": mode,
            "reference_label": reference_label,
            "mean_top_k_similarity": mean_top,
            "top_matches": top_matches,
            "display_hint": display_hint,
        }

    def format_few_shot_block(
        self,
        user_prompt: str,
        persona: str,
        *,
        ollama_base_url: str,
        top_k: int | None = None,
        max_example_chars: int = 4500,
    ) -> str:
        if not _success_store_enabled():
            return ""
        k = top_k if top_k is not None else _success_few_shot_k()
        if k <= 0:
            return ""

        entries = self.load_entries()
        if not entries:
            return ""

        query = (user_prompt or "").strip()
        base = ollama_base_url.rstrip("/")
        model = rag_embed_model()

        q_emb = _ollama_embed_batches(base, model, [query])
        idx_map: list[int] = []
        doc_embs: list[list[float]] = []

        for i, e in enumerate(entries):
            p = str(e.get("prompt") or "")
            vec = e.get("embedding")
            if not isinstance(vec, list) or len(vec) < 8:
                batch = _ollama_embed_batches(base, model, [p])
                vec = batch[0] if batch and len(batch) == 1 else None
            if vec is None:
                continue
            idx_map.append(i)
            doc_embs.append(vec)

        scored: list[tuple[float, dict[str, Any]]] = []
        if q_emb and len(q_emb) == 1 and doc_embs:
            for rank_i, sim in _cosine_rank(q_emb[0], doc_embs):
                ent = entries[idx_map[rank_i]]
                s = float(sim)
                if ent.get("persona") == persona:
                    s += 0.08
                scored.append((s, ent))
        else:
            prompts = [str(e.get("prompt") or "") for e in entries]
            for rank_i, sim in _lexical_rank(query, prompts):
                ent = entries[rank_i]
                s = float(sim)
                if ent.get("persona") == persona:
                    s += 5.0
                scored.append((s, ent))

        scored.sort(key=lambda x: x[0], reverse=True)
        picked = scored[:k]

        parts: list[str] = []
        for n, (sim, ent) in enumerate(picked, start=1):
            deck = ent.get("deck") or {}
            slides = deck.get("slides") or []
            ds = deck.get("design_specs")
            blob = {"slides": slides}
            if ds is not None:
                blob["design_specs"] = ds
            deck_json = json.dumps(blob, ensure_ascii=False, default=str)
            if len(deck_json) > max_example_chars:
                deck_json = deck_json[: max_example_chars - 24] + "\n… [truncated]"
            parts.append(
                f"### Prior success #{n} (retrieval score {sim:.3f})\n"
                f"- **Historical strategic directive:** {ent.get('prompt', '')}\n"
                f"- **Persona used:** {ent.get('persona', '')}\n"
                f"- **Materialized slide JSON (mimic structure and depth, adapt to current context):**\n"
                f"```json\n{deck_json}\n```",
            )

        return (
            "## Strategic memory — similar past materializations\n\n"
            "Use these as **few-shot** references for slide count, tone, and JSON shape; "
            "do not copy verbatim if the current directive or ingested files differ.\n\n"
            + "\n\n".join(parts)
        )

